"""Read judges' Bookings + Items sheets (tolerant to header row position, naming, Excel quirks)."""
from __future__ import annotations

import io
import re
from collections import OrderedDict

import openpyxl

from .models import (Booking, Item, to_num, to_text, to_hs, to_date, to_bool, norm_dest_type,
                     norm_priority, norm_city, norm_country, dest_country_for, is_blank)
from .reference import ORIGIN_CITIES, DEST_CITIES, DEST_COUNTRIES

BOOKING_COLS = {
    "bookingid": "booking_id", "sellername": "seller_name", "origincity": "origin_city",
    "origincountry": "origin_country", "destinationcity": "dest_city", "destinationcountry": "dest_country",
    "destinationtype": "dest_type", "readytoshipdate": "ready_date", "sellerpriority": "priority",
    "pallets": "pallets", "cartons": "cartons", "totaldeclaredvalue": "declared_value",
    "currency": "currency", "totalgrossweightkg": "gross_weight", "totalgrossweight": "gross_weight",
    "deliveryappointmentbooked": "appointment",
}
ITEM_COLS = {
    "bookingid": "booking_id", "sku": "sku", "description": "description", "hscode": "hs_code",
    "quantityunits": "quantity", "quantity": "quantity", "unitprice": "unit_price",
    "unitspercarton": "units_per_carton", "cartons": "cartons", "unitweightkg": "unit_weight",
    "unitweight": "unit_weight", "countryoforigin": "origin_country",
}
# The grey example rows that ship with the template (skipped if left in the file)
EXAMPLE_BOOKING = ("BK-001", "Northfield Goods", "Manchester", "Frankfurt", 3.0, 90.0, 5400.0, 620.0)
EXAMPLE_ITEM = ("BK-001", "MUG-01", "Ceramic coffee mug, 350 ml", 1080.0, 5.0, 12.0, 90.0)


def _key(s) -> str:
    return re.sub(r"[^a-z]", "", str(s or "").lower())


def _read_table(ws, colmap):
    rows = list(ws.iter_rows(values_only=True))
    header_idx = None
    for i, r in enumerate(rows[:15]):
        keys = [_key(c) for c in r]
        if "bookingid" in keys:
            header_idx = i
            break
    if header_idx is None:
        return None, []
    header = [colmap.get(_key(c)) for c in rows[header_idx]]
    out = []
    for n, r in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        rec = {}
        for h, v in zip(header, r):
            if h and h not in rec:
                rec[h] = v
        if all(is_blank(v) for v in rec.values()):
            continue
        rec["_row"] = n
        out.append(rec)
    return header, out


def _classify(ws):
    """Return 'bookings' / 'items' / None based on headers."""
    for r in list(ws.iter_rows(values_only=True))[:15]:
        keys = {_key(c) for c in r}
        if "bookingid" in keys:
            if "sku" in keys or "hscode" in keys or "unitprice" in keys:
                return "items"
            if "pallets" in keys or "sellername" in keys:
                return "bookings"
    return None


def load_workbooks(files: list[bytes]):
    """Accepts one workbook with both sheets, or two separate files. Returns (bookings, notes)."""
    sheets = {"bookings": None, "items": None}
    for data in files:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        for ws in wb.worksheets:
            kind = _classify(ws)
            if kind and sheets[kind] is None:
                sheets[kind] = ws
    notes = []
    if sheets["bookings"] is None:
        raise ValueError("Couldn't find a Bookings sheet (a sheet with 'Booking ID' and 'Pallets' columns).")
    _, brows = _read_table(sheets["bookings"], BOOKING_COLS)
    irows = []
    if sheets["items"] is not None:
        _, irows = _read_table(sheets["items"], ITEM_COLS)
    else:
        notes.append("No Items sheet found, so item-level checks can't run.")

    bookings: "OrderedDict[str, Booking]" = OrderedDict()
    for r in brows:
        bid = to_text(r.get("booking_id")) or f"(row {r['_row']})"
        sig = (bid, to_text(r.get("seller_name")), to_text(r.get("origin_city")), to_text(r.get("dest_city")),
               to_num(r.get("pallets")), to_num(r.get("cartons")), to_num(r.get("declared_value")),
               to_num(r.get("gross_weight")))
        if sig == EXAMPLE_BOOKING and r["_row"] == 3:
            notes.append("Skipped the grey template example row (row 3, BK-001 Northfield Goods).")
            continue
        if bid in bookings:
            notes.append(f"Booking ID {bid} appears more than once on the Bookings sheet; using the first row.")
            continue
        dest_city = norm_city(r.get("dest_city"), list(DEST_CITIES))
        bookings[bid] = Booking(
            booking_id=bid,
            seller_name=to_text(r.get("seller_name")),
            origin_city=norm_city(r.get("origin_city"), ORIGIN_CITIES),
            origin_country=norm_country(r.get("origin_country"), ["United Kingdom"]) or "United Kingdom",
            dest_city=dest_city,
            dest_country=norm_country(r.get("dest_country"), DEST_COUNTRIES) or dest_country_for(dest_city),
            dest_type=norm_dest_type(r.get("dest_type")),
            ready_date=to_date(r.get("ready_date")),
            priority=norm_priority(r.get("priority")),
            pallets=to_num(r.get("pallets")),
            cartons=to_num(r.get("cartons")),
            declared_value=to_num(r.get("declared_value")),
            currency=to_text(r.get("currency")) or "GBP",
            gross_weight=to_num(r.get("gross_weight")),
            appointment=to_bool(r.get("appointment")),
        )

    skipped_example_item = any("template example" in n for n in notes)
    orphans = set()
    for r in irows:
        bid = to_text(r.get("booking_id"))
        sig = (bid, to_text(r.get("sku")), to_text(r.get("description")), to_num(r.get("quantity")),
               to_num(r.get("unit_price")), to_num(r.get("units_per_carton")), to_num(r.get("cartons")))
        if skipped_example_item and sig == EXAMPLE_ITEM and r["_row"] == 3:
            continue
        it = Item(
            booking_id=bid or "",
            sku=to_text(r.get("sku")) or "",
            description=to_text(r.get("description")) or "",
            hs_code=to_hs(r.get("hs_code")),
            quantity=to_num(r.get("quantity")),
            unit_price=to_num(r.get("unit_price")),
            units_per_carton=to_num(r.get("units_per_carton")),
            cartons=to_num(r.get("cartons")),
            unit_weight=to_num(r.get("unit_weight")),
            origin_country=norm_country(r.get("origin_country")),
        )
        if bid in bookings:
            bookings[bid].items.append(it)
        else:
            orphans.add(bid or "(blank)")
    if orphans:
        notes.append(f"Item rows reference booking IDs not on the Bookings sheet: {', '.join(sorted(orphans))}. "
                     f"They were not checked.")
    return list(bookings.values()), notes
