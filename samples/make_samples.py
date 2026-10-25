"""Builds samples/sample_test_bookings.xlsx from the official template, with planted errors.
Run:  python samples/make_samples.py"""
from datetime import datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
wb = openpyxl.load_workbook(ROOT / "data" / "template.xlsx")
B, I = wb["Bookings"], wb["Items"]

# Booking ID, Seller, Origin city, Origin ctry, Dest city, Dest ctry, Dest type, Ready, Priority,
# Pallets, Cartons, Declared value, Currency, Gross kg, Appointment
bookings = [
    ("BK-101", "Thames Homeware", "London", "United Kingdom", "Berlin", "Germany", "Warehouse", "10-11-2026", "Cheapest", 2, 60, 3600, "GBP", 520, "Yes"),            # clean
    ("BK-102", "Pennine Crafts", "Leeds", "United Kingdom", "Paris", "France", "Business address", "12-11-2026", "Fastest", 1, 20, 900, "GBP", 180, "No"),            # R-01, R-09
    ("BK-103", "Brum Textiles", "Birmingham", "United Kingdom", "Amsterdam", "Netherlands", "Warehouse", "09-11-2026", "Most reliable", 2, 80, 4000, "GBP", 600, "Yes"),  # R-02
    ("BK-104", "Mancunian Toys", "Manchester", "United Kingdom", "Hamburg", "Germany", "Warehouse", "11-11-2026", "Balanced", 3, 95, 6720, "GBP", 900, "Yes"),        # R-03, R-04
    ("BK-105", "Heavy Metal Ltd", "Leeds", "United Kingdom", "Lyon", "France", "Business address", "16-11-2026", "Cheapest", 1, 40, 2400, "GBP", 820, "No"),          # R-05 (all carriers)
    ("BK-106", "Gadget Garage", "London", "United Kingdom", "Rotterdam", "Netherlands", "Warehouse", "13-11-2026", "Balanced", 1, 30, 4050, "GBP", 160, "Yes"),        # R-06
    ("BK-107", "Northern Paper Co", "Manchester", "United Kingdom", "Frankfurt", "Germany", "Warehouse", "18-11-2026", "Cheapest", 2, 100, 2500, "GBP", 300, "No"),    # R-07, R-08
    ("BK-108", "Old Date Traders", "Birmingham", "United Kingdom", "Paris", "France", "Business address", datetime(2026, 10, 1), "Balanced", 1, 25, 1250, "GBP", 150, "No"),  # R-10
    ("BK-109", "Speedy Spares", "London", "United Kingdom", "Berlin", "Germany", "Warehouse", "20-11-2026", "Fastest", 2, 60, 6000, "GBP", 900, "Yes"),              # R-05 amber note: C excluded by weight, booked on B
    ("BK-110", "Bean There Coffee", "Leeds", "United Kingdom", "Amsterdam", "Netherlands", "Business address", "19-11-2026", "Most reliable", 1, 50, 3000, "GBP", 560, "No"),  # 5-digit HS; R-05 amber note
]
# Booking ID, SKU, Description, HS, Qty, Unit price, Units/carton, Cartons, Unit kg, Origin
items = [
    ("BK-101", "MUG-02", "Stoneware mug, 400 ml", 691200, 720, 3.5, 24, 30, 0.45, "China"),
    ("BK-101", "BWL-01", "Ceramic serving bowl", 691200, 360, 3.0, 12, 30, 0.5, "United Kingdom"),
    ("BK-102", "QLT-01", "Handmade cotton quilt", None, 20, 45, 1, 20, 6, None),
    ("BK-103", "TWL-01", "Cotton bath towel", 630260, 800, 5.5, 10, 80, 0.6, "India"),
    ("BK-104", "BLK-01", "Wooden building blocks set", 950300, 1200, 4.2, 24, 45, 0.4, "Vietnam"),
    ("BK-104", "PZL-01", "Jigsaw puzzle 1000 pieces", 950590, 480, 3.5, 12, 40, 0.6, "United Kingdom"),
    ("BK-105", "PLT-01", "Cast iron weight plate 20 kg", 950691, 40, 60, 1, 40, 20, "China"),
    ("BK-106", "PB-10K", "Portable power bank 10000mAh lithium-ion", 850760, 300, 9.5, 20, 15, 0.25, "China"),
    ("BK-106", "PRF-50", "Eau de parfum 50 ml", 330300, 150, 8.0, 10, 15, 0.2, "Turkey"),
    ("BK-107", "NB-A5", "A5 lined notebook", 482020, 5000, 0.5, 50, 100, 0.12, "United Kingdom"),
    ("BK-108", "CNDL-1", "Soy wax candle in glass jar", 340600, 500, 2.5, 20, 25, 0.25, "United Kingdom"),
    ("BK-109", "BRK-PAD", "Brake pads set", 870830, 1200, 5.0, 20, 60, 0.6, "Germany"),
    ("BK-110", "COF-1K", "Roasted coffee beans 1 kg bag", 90121, 500, 6.0, 10, 50, 1.0, "Vietnam"),
]
for r, row in enumerate(bookings, start=4):
    for c, v in enumerate(row, start=1):
        B.cell(r, c, v)
for r, row in enumerate(items, start=4):
    for c, v in enumerate(row, start=1):
        I.cell(r, c, v)

# judges only upload Sheets 1 & 2 -> drop the rest to mimic that
for name in wb.sheetnames[2:]:
    del wb[name]
out = ROOT / "samples" / "sample_test_bookings.xlsx"
wb.save(out)
print("wrote", out)
